import csv
import json
import re
from collections.abc import Callable, Iterable
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from docx import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph
from lxml.etree import XMLSyntaxError
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pypdf.errors import PdfReadError

from app.core.document_formats import DOCUMENT_FORMATS, document_extension

CHUNK_SIZE_TOKENS = 600
CHUNK_OVERLAP_TOKENS = 80

_splitter = RecursiveCharacterTextSplitter.from_tiktoken_encoder(
    encoding_name="cl100k_base",
    chunk_size=CHUNK_SIZE_TOKENS,
    chunk_overlap=CHUNK_OVERLAP_TOKENS,
)


class DocumentParsingError(ValueError):
    pass


class UnsupportedDocumentFormatError(DocumentParsingError):
    pass


class EmptyDocumentError(DocumentParsingError):
    pass


def supported_document_extensions() -> tuple[str, ...]:
    return tuple(_EXTRACTORS)


def parse_document(file_path: Path) -> list[Document]:
    extension = document_extension(file_path.name)
    extractor = _EXTRACTORS.get(extension)
    if extractor is None:
        raise UnsupportedDocumentFormatError(
            f"Unsupported document format '{extension or 'unknown'}'."
        )

    try:
        elements = extractor(file_path)
    except (
        BadZipFile,
        csv.Error,
        DocxPackageNotFoundError,
        InvalidFileException,
        json.JSONDecodeError,
        KeyError,
        ParseError,
        PdfReadError,
        PptxPackageNotFoundError,
        UnicodeDecodeError,
        XMLSyntaxError,
    ) as error:
        raise DocumentParsingError(
            f"Could not parse the uploaded {extension} document."
        ) from error
    chunks = _chunk_elements(elements, extension)
    if not chunks:
        raise EmptyDocumentError(
            "No extractable text was found. Scanned PDFs require OCR, which is not enabled."
            if extension == ".pdf"
            else "No extractable text was found in the document."
        )
    return chunks


def _read_text(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8-sig")


def _extract_text(file_path: Path) -> list[Document]:
    text = _read_text(file_path)
    return [Document(page_content=part) for part in re.split(r"\n\s*\n", text) if part.strip()]


def _extract_markdown(file_path: Path) -> list[Document]:
    elements: list[Document] = []
    section_path: list[str] = []
    current_lines: list[str] = []

    def flush() -> None:
        text = "\n".join(current_lines).strip()
        if text:
            elements.append(Document(
                page_content=text,
                metadata={"section": " / ".join(section_path)} if section_path else {},
            ))
        current_lines.clear()

    for line in _read_text(file_path).splitlines():
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            section_path[:] = section_path[:level - 1]
            section_path.append(title)
            current_lines.append(line)
        elif line.strip():
            current_lines.append(line)
        elif current_lines and current_lines[-1]:
            current_lines.append("")
    flush()
    return elements


def _column_name(headers: list[str], index: int) -> str:
    if index < len(headers) and headers[index]:
        return headers[index]
    return f"Column {index + 1}"


def _extract_docx(file_path: Path) -> list[Document]:
    source = WordDocument(file_path)
    elements: list[Document] = []
    section_path: list[str] = []
    table_index = 0
    for child in source.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, source)
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            heading = re.fullmatch(r"Heading ([1-9])", style_name)
            if heading:
                level = int(heading.group(1))
                section_path[:] = section_path[:level - 1]
                section_path.append(text)
            elements.append(Document(
                page_content=text,
                metadata={
                    "section": " / ".join(section_path),
                    "content_type": "heading" if heading else "paragraph",
                },
            ))
        elif child.tag == qn("w:tbl"):
            table = WordTable(child, source)
            rows = [
                [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                for row in table.rows
            ]
            if rows:
                headers = rows[0]
                for row_index, row in enumerate(rows[1:] or rows):
                    values = []
                    for index, value in enumerate(row):
                        if value:
                            values.append(f"{_column_name(headers, index)}: {value}")
                    if values:
                        elements.append(Document(
                            page_content=" | ".join(values),
                            metadata={
                                "section": " / ".join(section_path),
                                "content_type": "table",
                                "table_index": table_index,
                                "row_index": row_index,
                            },
                        ))
            table_index += 1
    return elements


def _extract_pptx(file_path: Path) -> list[Document]:
    presentation = Presentation(file_path)
    elements: list[Document] = []
    for slide_index, slide in enumerate(presentation.slides):
        title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        texts = [title] if title else []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame and shape.text.strip():
                texts.append(shape.text.strip())
            if shape.has_table:
                table = shape.table
                rows = [
                    [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                    for row in table.rows
                ]
                if rows:
                    headers = rows[0]
                    for row in rows[1:]:
                        values = []
                        for index, value in enumerate(row):
                            if value:
                                values.append(f"{_column_name(headers, index)}: {value}")
                        if values:
                            texts.append(" | ".join(values))
        if texts:
            elements.append(Document(
                page_content="\n\n".join(texts),
                metadata={
                    "page": slide_index,
                    "section": title,
                    "content_type": "slide",
                },
            ))
    return elements


def _display_value(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _row_element(
    headers: list[str],
    values: Iterable[Any],
    row_index: int,
    section: str | None = None,
) -> Document | None:
    pairs = [
        f"{_column_name(headers, index)}: {_display_value(value)}"
        for index, value in enumerate(values)
        if _display_value(value)
    ]
    if not pairs:
        return None
    metadata: dict[str, str | int] = {
        "content_type": "table_row",
        "row_index": row_index,
    }
    if section:
        metadata["sheet"] = section
    return Document(page_content=" | ".join(pairs), metadata=metadata)


def _extract_csv(file_path: Path) -> list[Document]:
    elements: list[Document] = []
    with file_path.open("r", encoding="utf-8-sig", newline="") as source:
        delimiter = "\t" if document_extension(file_path.name) == ".tsv" else ","
        rows = csv.reader(source, delimiter=delimiter)
        headers = next(rows, [])
        for row_index, row in enumerate(rows, start=2):
            element = _row_element(headers, row, row_index)
            if element is not None:
                elements.append(element)
    return elements


def _extract_xlsx(file_path: Path) -> list[Document]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    elements: list[Document] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [_display_value(value) for value in next(rows, ())]
            for row_index, row in enumerate(rows, start=2):
                element = _row_element(headers, row, row_index, sheet.title)
                if element is not None:
                    elements.append(element)
    finally:
        workbook.close()
    return elements


class _VisibleTextExtractor(HTMLParser):
    _ignored_tags = {"script", "style", "noscript", "svg"}
    _block_tags = {
        "article", "blockquote", "br", "div", "h1", "h2", "h3",
        "h4", "h5", "h6", "li", "p", "tr",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.elements: list[Document] = []
        self._parts: list[str] = []
        self._section_path: list[str] = []
        self._heading_level: int | None = None
        self._ignored_depth = 0
        self._in_table = False
        self._in_table_row = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._ignored_tags:
            self._ignored_depth += 1
            return
        if self._ignored_depth == 0:
            if tag in self._block_tags:
                self._flush()
                if tag.startswith("h") and tag[1:].isdigit():
                    self._heading_level = int(tag[1:])
                if tag == "tr":
                    self._in_table_row = True
            elif tag == "table":
                self._in_table = True
            elif tag in {"td", "th"}:
                if self._parts:
                    self._parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._ignored_tags and self._ignored_depth:
            self._ignored_depth -= 1
            return
        if self._ignored_depth == 0:
            if tag in self._block_tags:
                self._flush()
                if tag.startswith("h") and tag[1:].isdigit():
                    self._heading_level = None
                if tag == "tr":
                    self._in_table_row = False
            elif tag == "table":
                self._in_table = False

    def handle_data(self, data: str) -> None:
        if self._ignored_depth == 0:
            self._parts.append(data)

    def _flush(self) -> None:
        text = re.sub(r"\s+", " ", "".join(self._parts)).strip()
        self._parts.clear()
        if not text:
            return
        if self._heading_level is not None:
            self._section_path[:] = self._section_path[:self._heading_level - 1]
            self._section_path.append(text)
            content_type = "heading"
        else:
            content_type = "table_row" if self._in_table_row or self._in_table else "paragraph"
        metadata: dict[str, str] = {"content_type": content_type}
        if self._section_path:
            metadata["section"] = " / ".join(self._section_path)
        self.elements.append(Document(page_content=text, metadata=metadata))


def _extract_html(file_path: Path) -> list[Document]:
    parser = _VisibleTextExtractor()
    parser.feed(_read_text(file_path))
    parser.close()
    parser._flush()
    return parser.elements


def _extract_pdf(file_path: Path) -> list[Document]:
    return PyPDFLoader(file_path).load()


def _extract_json(file_path: Path) -> list[Document]:
    value = json.loads(_read_text(file_path))
    return [Document(page_content=json.dumps(value, indent=2))]


_DocumentExtractor = Callable[[Path], list[Document]]
_EXTRACTORS: dict[str, _DocumentExtractor] = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".txt": _extract_text,
    ".md": _extract_markdown,
    ".markdown": _extract_markdown,
    ".csv": _extract_csv,
    ".tsv": _extract_csv,
    ".json": _extract_json,
    ".html": _extract_html,
    ".htm": _extract_html,
}
if _EXTRACTORS.keys() != DOCUMENT_FORMATS.keys():
    raise RuntimeError("Document upload formats and parser registrations must match.")


def _chunk_elements(elements: list[Document], extension: str) -> list[Document]:
    spans: list[tuple[int, int, dict[str, Any]]] = []
    text_parts: list[str] = []
    offset = 0
    separator = "\n" if extension == ".pdf" else "\n\n"
    for element in elements:
        content = element.page_content.strip()
        if not content:
            continue
        if text_parts:
            text_parts.append(separator)
            offset += len(separator)
        start = offset
        text_parts.append(content)
        offset += len(content)
        spans.append((start, offset, element.metadata or {}))

    combined_text = "".join(text_parts).strip()
    if not combined_text:
        return []

    split_documents = _splitter.split_text(combined_text)
    chunks: list[Document] = []
    search_from = 0
    for content in split_documents:
        start = combined_text.find(content, search_from)
        if start < 0:
            raise RuntimeError("The token-aware splitter did not return a valid chunk start index.")
        end = start + len(content)
        search_from = start + 1
        relevant_metadata = [
            metadata for element_start, element_end, metadata in spans
            if element_end > start and element_start < end
        ]
        metadata = _chunk_metadata(relevant_metadata, extension)
        chunks.append(Document(page_content=content, metadata=metadata))
    return chunks


def _chunk_metadata(
    element_metadata: list[dict[str, Any]],
    extension: str,
) -> dict[str, Any]:
    metadata: dict[str, Any] = {"file_format": extension.lstrip(".")}
    pages = sorted({
        value
        for item in element_metadata
        if isinstance((value := item.get("page")), int) and not isinstance(value, bool)
    })
    if pages:
        metadata["page"] = pages[0]
        metadata["page_start"] = pages[0]
        metadata["page_end"] = pages[-1]

    sections = list(dict.fromkeys(
        section for item in element_metadata
        if isinstance((section := item.get("section")), str) and section
    ))
    if len(sections) == 1:
        metadata["section"] = sections[0]
    elif sections:
        metadata["section"] = " / ".join(sections)
        metadata["sections"] = sections

    content_types = list(dict.fromkeys(
        content_type for item in element_metadata
        if isinstance((content_type := item.get("content_type")), str) and content_type
    ))
    if len(content_types) == 1:
        metadata["content_type"] = content_types[0]
    elif content_types:
        metadata["content_type"] = "mixed"

    sheets = list(dict.fromkeys(
        sheet for item in element_metadata
        if isinstance((sheet := item.get("sheet")), str) and sheet
    ))
    if sheets:
        metadata["sheets"] = sheets
        if len(sheets) == 1:
            metadata["section"] = f"Sheet: {sheets[0]}"
    return metadata
