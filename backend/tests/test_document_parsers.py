import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document as WordDocument
from langchain_core.documents import Document
from openpyxl import Workbook
from pptx import Presentation
import tiktoken

from app.infrastructure.parsers.documents import (
    CHUNK_SIZE_TOKENS,
    EmptyDocumentError,
    UnsupportedDocumentFormatError,
    parse_document,
    supported_document_extensions,
)


class DocumentParserTests(unittest.TestCase):
    def test_supported_extensions_cover_common_document_formats(self):
        self.assertEqual(
            set(supported_document_extensions()),
            {
                ".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md",
                ".markdown", ".csv", ".json", ".html", ".htm",
                ".tsv",
            },
        )

    def test_markdown_chunks_preserve_heading_metadata(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.md"
            path.write_text("# Policies\n\nResidents may request access.", encoding="utf-8")

            chunks = parse_document(path)

        self.assertEqual(chunks[0].metadata["section"], "Policies")
        self.assertEqual(chunks[0].metadata["file_format"], "md")
        self.assertIn("Residents may request access.", chunks[0].page_content)

    def test_docx_extracts_heading_paragraphs_and_header_labeled_table_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.docx"
            source = WordDocument()
            source.add_heading("Access", level=1)
            source.add_paragraph("Use the lobby entrance.")
            table = source.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Item"
            table.cell(0, 1).text = "Location"
            table.cell(1, 0).text = "Key"
            table.cell(1, 1).text = "Front desk"
            source.save(path)

            chunks = parse_document(path)

        content = "\n".join(chunk.page_content for chunk in chunks)
        self.assertIn("Use the lobby entrance.", content)
        self.assertIn("Item: Key", content)
        self.assertIn("Location: Front desk", content)
        self.assertTrue(all(chunk.metadata["file_format"] == "docx" for chunk in chunks))
        self.assertTrue(any(chunk.metadata.get("section") == "Access" for chunk in chunks))

    def test_pptx_extracts_slide_text_and_citation_page(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "briefing.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = "Safety"
            slide.placeholders[1].text = "Call the service desk."
            presentation.save(path)

            chunks = parse_document(path)

        self.assertEqual(chunks[0].metadata["page"], 0)
        self.assertEqual(chunks[0].metadata["section"], "Safety")
        self.assertIn("Call the service desk.", chunks[0].page_content)

    def test_xlsx_repeats_column_labels_in_row_content(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "contacts.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Contacts"
            sheet.append(["Team", "Email"])
            sheet.append(["Support", "help@example.test"])
            workbook.save(path)

            chunks = parse_document(path)

        self.assertIn("Team: Support", chunks[0].page_content)
        self.assertIn("Email: help@example.test", chunks[0].page_content)
        self.assertEqual(chunks[0].metadata["section"], "Sheet: Contacts")

    def test_csv_rows_keep_header_context_and_json_is_extracted(self):
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "people.csv"
            with csv_path.open("w", encoding="utf-8", newline="") as output:
                writer = csv.writer(output)
                writer.writerow(["Name", "Role"])
                writer.writerow(["Ada", "Engineer"])
            json_path = Path(directory) / "settings.json"
            json_path.write_text(json.dumps({"retention_days": 30}), encoding="utf-8")

            csv_chunks = parse_document(csv_path)
            json_chunks = parse_document(json_path)

        self.assertIn("Name: Ada", csv_chunks[0].page_content)
        self.assertIn("Role: Engineer", csv_chunks[0].page_content)
        self.assertIn('"retention_days": 30', json_chunks[0].page_content)

    def test_html_omits_script_content_and_extracts_visible_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "page.html"
            path.write_text(
                "<h1>Getting started</h1><p>Read the <b>manual</b>.</p>"
                "<script>secret()</script>",
                encoding="utf-8",
            )

            chunks = parse_document(path)

        text = "\n".join(chunk.page_content for chunk in chunks)
        self.assertIn("Getting started", text)
        self.assertIn("Read the manual.", text)
        self.assertNotIn("secret", text)
        self.assertEqual(chunks[0].metadata["section"], "Getting started")

    def test_pdf_pages_are_combined_and_chunk_page_ranges_are_preserved(self):
        page_text = " ".join(f"term{index}" for index in range(120))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "long.pdf"
            path.touch()
            with patch(
                "app.infrastructure.parsers.documents.PyPDFLoader",
                return_value=type(
                    "Loader",
                    (),
                    {"load": lambda self: [
                        Document(page_content=page_text, metadata={"page": 0}),
                        Document(page_content=page_text, metadata={"page": 1}),
                    ]},
                )(),
            ):
                chunks = parse_document(path)

        encoding = tiktoken.get_encoding("cl100k_base")
        self.assertTrue(any(
            chunk.metadata.get("page_start") == 0
            and chunk.metadata.get("page_end") == 1
            for chunk in chunks
        ))
        self.assertTrue(all(len(encoding.encode(chunk.page_content)) <= CHUNK_SIZE_TOKENS for chunk in chunks))

    def test_unsupported_and_empty_documents_fail_explicitly(self):
        with tempfile.TemporaryDirectory() as directory:
            unsupported_path = Path(directory) / "archive.zip"
            unsupported_path.write_bytes(b"")
            empty_path = Path(directory) / "empty.txt"
            empty_path.write_text("", encoding="utf-8")

            with self.assertRaises(UnsupportedDocumentFormatError):
                parse_document(unsupported_path)
            with self.assertRaises(EmptyDocumentError):
                parse_document(empty_path)


if __name__ == "__main__":
    unittest.main()
